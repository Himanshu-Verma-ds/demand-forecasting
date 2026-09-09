from __future__ import annotations

"""Append a Markdown section to an EXISTING .docx, preserving everything already in it.

md_to_docx.py rebuilds the whole document from the Markdown source, which would discard any
edits made directly in Word. This opens the existing file, appends the rendered section at the
end, and saves in place - so manual changes survive.

    python scripts/append_docx_section.py final_conclusion/final_report.docx section.md
"""

import argparse
import shutil
from datetime import datetime
from pathlib import Path

from docx import Document

from md_to_docx import (  # noqa: E402  - same directory, imported for its renderers
    BLACK,
    BODY_FONT,
    add_code_block,
    add_inline,
    add_table,
    is_separator,
)
from docx.shared import Pt


def render_into(doc: Document, lines: list[str]) -> dict:
    """Render Markdown lines onto an already-open document, using md_to_docx's renderers."""
    import re

    stats = {"headings": 0, "tables": 0, "paragraphs": 0, "list_items": 0, "code_blocks": 0}
    i = 0

    while i < len(lines):
        stripped = lines[i].strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, block)
            stats["code_blocks"] += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            stats["tables"] += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if heading:
            level = min(len(heading.group(1)), 4)
            paragraph = doc.add_heading("", level=level)
            add_inline(paragraph, heading.group(2))

            sizes = {1: 17, 2: 13.5, 3: 11.5, 4: 10.5}
            paragraph.paragraph_format.space_before = Pt(16 if level <= 2 else 12)
            paragraph.paragraph_format.space_after = Pt(6)

            for run in paragraph.runs:
                run.font.name = BODY_FONT
                run.font.size = Pt(sizes[level])
                run.font.color.rgb = BLACK
                run.bold = True
                run.underline = True

            stats["headings"] += 1
            i += 1
            continue

        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            i += 1
            continue

        bullet = re.match(r"^[-*]\s+(.*)", stripped)
        numbered = re.match(r"^(\d+)\.\s+(.*)", stripped)

        if bullet or numbered:
            style = "List Bullet" if bullet else "List Number"
            paragraph = doc.add_paragraph(style=style)
            add_inline(paragraph, bullet.group(1) if bullet else numbered.group(2))
            for run in paragraph.runs:
                run.font.color.rgb = BLACK
            stats["list_items"] += 1
            i += 1
            continue

        block = []
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#{1,6}\s|```|\||>|[-*]\s|\d+\.\s|-{3,}$)", lines[i].strip()
        ):
            block.append(lines[i].strip())
            i += 1

        if block:
            paragraph = doc.add_paragraph()
            add_inline(paragraph, " ".join(block))
            stats["paragraphs"] += 1

    return stats


def main():
    ap = argparse.ArgumentParser(description="Append a Markdown section to an existing .docx.")
    ap.add_argument("docx")
    ap.add_argument("section_md")
    ap.add_argument("--no-backup", action="store_true")
    args = ap.parse_args()

    target = Path(args.docx)

    if not args.no_backup:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        backup = target.with_name(f"{target.stem}.backup-{stamp}{target.suffix}")
        shutil.copy2(target, backup)
        print(f"backed up existing document to {backup.name}")

    doc = Document(str(target))
    before = len(doc.paragraphs), len(doc.tables)

    stats = render_into(doc, Path(args.section_md).read_text(encoding="utf-8").splitlines())
    doc.save(str(target))

    after = len(Document(str(target)).paragraphs), len(Document(str(target)).tables)
    print(f"appended to {target}")
    print(f"  paragraphs {before[0]} -> {after[0]}, tables {before[1]} -> {after[1]}")
    for key, value in stats.items():
        print(f"  {key}: {value}")


if __name__ == "__main__":
    main()
